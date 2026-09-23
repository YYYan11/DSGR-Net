"""Streaming data preparation and cache builders for ICASSP experiments."""

from __future__ import annotations

import csv
import hashlib
import itertools
import json
from collections import defaultdict
from datetime import date
from pathlib import Path
from typing import Iterator

import numpy as np


WEATHER_FEATURES = ("SRAD", "TMAX", "TMIN", "RAIN", "WIND")
MANAGEMENT_FEATURES = ("IRR", "N", "P", "K")
SEQUENCE_FEATURES = WEATHER_FEATURES + MANAGEMENT_FEATURES

# Wang et al. (Agronomy, 2024) 的华兴农场田间试验协议。论文明确记载
# 中棉113于 2023-04-10 机播；迁移到 2024 外部验证时沿用相同月/日，
# 但不能把它表述为 2024 年现场实测播种日。
PAPER_FIELD_PROTOCOL = {
    "cultivar": "Zhongmian 113",
    "sowing_month": 4,
    "sowing_day": 10,
    "source_experiment_year": 2023,
    "emergence_irrigation_date": "2023-04-20",
    "cultivar_calibration": "16 historical phenology records; GLUE; 20,000 iterations",
    "dssat_validation": "2023 field yield and LAI",
    "reference": "Wang et al., Agronomy 2024, 14, 14; Wang et al., Agricultural Water Management 2025, 317, 109624",
}


def stream_json_array(path: str | Path, chunk_size: int = 1 << 20) -> Iterator[dict]:
    """逐项读取顶层 JSON 数组，避免一次载入数 GB 数据。"""
    decoder = json.JSONDecoder()
    with Path(path).open("r", encoding="utf-8") as handle:
        buffer = ""
        started = False
        eof = False
        while True:
            if not eof and len(buffer) < chunk_size:
                chunk = handle.read(chunk_size)
                eof = not chunk
                buffer += chunk
            position = 0
            while position < len(buffer) and buffer[position].isspace():
                position += 1
            if not started:
                if position >= len(buffer):
                    if eof:
                        return
                    buffer = ""
                    continue
                if buffer[position] != "[":
                    raise ValueError("JSON 顶层必须是数组")
                position += 1
                started = True
            while position < len(buffer) and (
                buffer[position].isspace() or buffer[position] == ","
            ):
                position += 1
            if position < len(buffer) and buffer[position] == "]":
                return
            if position >= len(buffer):
                if eof:
                    raise ValueError("JSON 数组意外结束")
                buffer = ""
                continue
            try:
                item, end = decoder.raw_decode(buffer, position)
            except json.JSONDecodeError:
                if eof:
                    raise
                buffer = buffer[position:]
                continue
            if not isinstance(item, dict):
                raise TypeError("每条训练记录必须是对象")
            yield item
            buffer = buffer[end:]


def _natural_sort_key(value: str) -> tuple:
    text = str(value)
    prefix, _, suffix = text.rpartition("_")
    return (prefix or text, int(suffix) if suffix.isdigit() else -1, text)


def _build_static_template(record: dict) -> list[tuple[str, ...]]:
    template: list[tuple[str, ...]] = []
    for field in ("SH2O", "SNH4", "SNO3"):
        for key in sorted(record[field], key=lambda value: int(value)):
            template.append(("record", field, key))
    for field in (
        "Plant_time(How_many_days_of_the_year)",
        "Seeding_density(p/m2)",
        "Seeding_row_spacing(cm)",
        "Seeding_depth(cm)",
    ):
        template.append(("record", field))
    soil = record["Soil_data"]
    for key in sorted(soil, key=_natural_sort_key):
        value = soil[key]
        if isinstance(value, dict):
            for subkey in sorted(value):
                template.append(("soil", key, subkey))
        else:
            template.append(("soil", key))
    weather = record["wthear_data"]
    for key in sorted(key for key, value in weather.items() if not isinstance(value, dict)):
        template.append(("weather", key))
    return template


def _get_value(record: dict, path: tuple[str, ...]) -> float:
    scope, *keys = path
    value = record if scope == "record" else record["Soil_data" if scope == "soil" else "wthear_data"]
    for key in keys:
        value = value[key]
    return float(value)


def _environment_id(record: dict) -> str:
    environment = {
        key: record[key]
        for key in (
            "SH2O",
            "SNH4",
            "SNO3",
            "Plant_time(How_many_days_of_the_year)",
            "Seeding_density(p/m2)",
            "Seeding_row_spacing(cm)",
            "Seeding_depth(cm)",
            "Soil_data",
            "wthear_data",
        )
    }
    payload = json.dumps(
        environment, sort_keys=True, ensure_ascii=False, separators=(",", ":")
    ).encode()
    return hashlib.sha256(payload).hexdigest()


def _management_sequence(record: dict, start_day: int, end_day: int) -> np.ndarray:
    days = end_day - start_day + 1
    if SEQUENCE_FEATURES != WEATHER_FEATURES + MANAGEMENT_FEATURES:
        raise AssertionError("sequence feature order is inconsistent")
    sequence = np.zeros((days, len(SEQUENCE_FEATURES)), dtype=np.float32)
    weather = record["wthear_data"]
    for offset, day in enumerate(range(start_day, end_day + 1)):
        daily = weather.get(str(day))
        if not isinstance(daily, dict):
            raise ValueError(f"缺少第 {day} 天气象记录")
        sequence[offset, :len(WEATHER_FEATURES)] = [float(daily[name]) for name in WEATHER_FEATURES]
    for column, field in enumerate(
        (
            "irrigattons(mm)",
            "fertilizers_N(kg/ha)",
            "fertilizers_P(kg/ha)",
            "fertilizers_K(kg/ha)",
        ),
        start=5,
    ):
        for date, value in record[field].items():
            text = str(date)
            if len(text) != 5 or not text.isdigit():
                raise ValueError(f"尚未支持的管理日期编码：{date}")
            day = int(text[-3:])
            if not 1 <= day <= 366:
                raise ValueError(f"年积日超出范围：{date}")
            if start_day <= day <= end_day:
                sequence[day - start_day, column] += float(value)
    return sequence


def _wl_feature_vector(record: dict) -> np.ndarray:
    """Reproduce the feature order used by Wang's original WL-DNN scripts.

    The legacy code concatenates twelve management slots, three ten-layer
    initial soil states, four planting descriptors, nested soil fields in JSON
    insertion order, and 365 daily weather records.  A leap-year record drops
    February 29 so that calendar dates retain the same 365 input positions.
    """
    values: list[float] = []
    for field in (
        "irrigattons(mm)",
        "fertilizers_N(kg/ha)",
        "fertilizers_P(kg/ha)",
        "fertilizers_K(kg/ha)",
    ):
        series = record[field]
        if len(series) != 12:
            raise ValueError(f"WL-DNN expects 12 management slots for {field}; got {len(series)}")
        values.extend(float(value) for value in series.values())
    for field in ("SH2O", "SNH4", "SNO3"):
        if len(record[field]) != 10:
            raise ValueError(f"WL-DNN expects 10 soil layers for {field}; got {len(record[field])}")
        values.extend(float(value) for value in record[field].values())
    for field in (
        "Plant_time(How_many_days_of_the_year)",
        "Seeding_density(p/m2)",
        "Seeding_row_spacing(cm)",
        "Seeding_depth(cm)",
    ):
        values.append(float(record[field]))
    for value in record["Soil_data"].values():
        if isinstance(value, dict):
            values.extend(float(item) for item in value.values())

    daily_weather = [
        (str(day), value)
        for day, value in record["wthear_data"].items()
        if isinstance(value, dict)
    ]
    if len(daily_weather) == 366:
        daily_weather = [(day, value) for day, value in daily_weather if int(day) != 60]
    if len(daily_weather) != 365:
        raise ValueError(f"WL-DNN expects 365 weather days; got {len(daily_weather)}")
    for _, value in daily_weather:
        values.extend(float(item) for item in value.values())
    result = np.asarray(values, dtype=np.float32)
    if result.shape != (1997,):
        raise ValueError(f"WL-DNN feature dimension must be 1997; got {result.shape[0]}")
    return result


def _file_sha256(path: str | Path, chunk_size: int = 1 << 20) -> str:
    digest = hashlib.sha256()
    with Path(path).open("rb") as handle:
        while block := handle.read(chunk_size):
            digest.update(block)
    return digest.hexdigest()


def _to_float(value: object) -> float:
    if value in (None, "", "-"):
        return 0.0
    return float(value)


def _read_2024_weather(path: str | Path) -> tuple[dict[str, dict[str, float]], float, float]:
    daily: dict[str, dict[str, float]] = {}
    monthly_temperature: dict[int, list[float]] = defaultdict(list)
    with Path(path).open("r", encoding="utf-8-sig", errors="replace", newline="") as handle:
        reader = csv.DictReader(handle)
        header = tuple((name or "").strip() for name in (reader.fieldnames or ()))
        if not header:
            raise ValueError("气象 CSV 缺少表头")
        reader.fieldnames = list(header)
        normalized = {name.lower(): name for name in header}
        for required in ("date", "rain(mm/day)", "wind(km/d)"):
            if required not in normalized:
                raise ValueError(f"气象 CSV 缺少必需列：{required}；实际列为 {header}")

        def find_unique_prefix(prefix: str) -> str:
            candidates = [name for name in header if name.lower().startswith(prefix)]
            if len(candidates) != 1:
                raise ValueError(f"气象 CSV 的 {prefix} 列应唯一匹配，实际为 {candidates}")
            return candidates[0]

        date_column = normalized["date"]
        rain_column = normalized["rain(mm/day)"]
        wind_column = normalized["wind(km/d)"]
        tmax_column = find_unique_prefix("tmax")
        tmin_column = find_unique_prefix("tmin")
        srad_column = find_unique_prefix("srad")
        for row in reader:
            date_value = str(row.get(date_column, "")).strip()
            if not date_value.startswith("2024-"):
                continue
            year, month, day_of_month = (int(part) for part in date_value.split("-"))
            current = date(year, month, day_of_month)
            day = str(current.timetuple().tm_yday)
            tmax = float(row[tmax_column])
            tmin = float(row[tmin_column])
            daily[day] = {
                "SRAD": float(row[srad_column]),
                "TMAX": tmax,
                "TMIN": tmin,
                "RAIN": float(row[rain_column]),
                "WIND": float(row[wind_column]) / 86.4,
            }
            monthly_temperature[current.month].append((tmax + tmin) / 2.0)
    if len(daily) != 366:
        raise ValueError(f"2024 年气象应有 366 天，实际 {len(daily)} 天")
    monthly_mean = [np.mean(monthly_temperature[month]) for month in range(1, 13)]
    tav = float(np.mean([(row["TMAX"] + row["TMIN"]) / 2.0 for row in daily.values()]))
    amp = float(max(monthly_mean) - min(monthly_mean))
    return daily, tav, amp


def read_2024_sensor_soil_moisture(
    workbook: str | Path,
    treatments: tuple[str, ...] = ("M1", "M2", "M3", "M4", "M5", "M6"),
) -> dict:
    """读取 2024 地块土壤湿度，按传感器日中位数和地块中位数稳健聚合。

    传感器表没有给出可与根区有效水量直接换算的深度/标定关系，因此这里只
    保留 0--100% 内的相对湿度轨迹，用于验证逐日变化方向，不做绝对 mm 拟合。
    """
    from openpyxl import load_workbook

    workbook = Path(workbook)
    book = load_workbook(workbook, read_only=True, data_only=True)

    def normalize_model_number(value: object) -> str:
        return str(value).strip().strip("'")

    model_humidity_sensors: dict[str, set[int]] = defaultdict(set)
    for row in book["Sensors"].iter_rows(min_row=3, values_only=True):
        sensor_id, sensor_type, model_number = row[1], row[2], row[3]
        if (
            sensor_id not in (None, "")
            and str(sensor_type).strip() == "HumiditySensor"
            and model_number not in (None, "")
        ):
            model_humidity_sensors[normalize_model_number(model_number)].add(int(sensor_id))

    land_models: dict[str, list[str]] = defaultdict(list)
    current_land: str | None = None
    for land, model_number in book["Land"].iter_rows(min_row=3, values_only=True):
        if land not in (None, ""):
            current_land = str(land).strip()
        if current_land and model_number not in (None, ""):
            land_models[current_land].append(normalize_model_number(model_number))

    treatment_sensors = {
        treatment: sorted({
            sensor_id
            for model_number in land_models.get(treatment, ())
            for sensor_id in model_humidity_sensors.get(model_number, ())
        })
        for treatment in treatments
    }
    sensor_treatments = {
        sensor_id: treatment
        for treatment, sensor_ids in treatment_sensors.items()
        for sensor_id in sensor_ids
    }
    if not sensor_treatments:
        book.close()
        raise ValueError("Land/Sensors 表未找到 M1--M6 的土壤湿度传感器映射")

    sensor_daily: dict[tuple[str, int, date], list[float]] = defaultdict(list)
    rejected_values = 0
    for row in book["HumiditySensor"].iter_rows(min_row=3, values_only=True):
        sensor_id, work_status, collected_at, value = row[1], row[2], row[6], row[7]
        if sensor_id in (None, "") or int(sensor_id) not in sensor_treatments:
            continue
        if int(work_status or 0) != 1 or collected_at is None or value in (None, ""):
            continue
        try:
            humidity = float(value)
        except (TypeError, ValueError):
            rejected_values += 1
            continue
        if not np.isfinite(humidity) or not 0.0 < humidity <= 100.0:
            rejected_values += 1
            continue
        current_date = collected_at.date() if hasattr(collected_at, "date") else date.fromisoformat(str(collected_at)[:10])
        treatment = sensor_treatments[int(sensor_id)]
        sensor_daily[(treatment, int(sensor_id), current_date)].append(humidity)
    book.close()

    land_daily: dict[tuple[str, date], list[float]] = defaultdict(list)
    for (treatment, _, current_date), values in sensor_daily.items():
        land_daily[(treatment, current_date)].append(float(np.median(values)))
    daily: dict[str, dict[str, float]] = defaultdict(dict)
    for (treatment, current_date), values in land_daily.items():
        if current_date.year == 2024:
            daily[treatment][str(current_date.timetuple().tm_yday)] = float(np.median(values))

    coverage = {}
    for treatment in treatments:
        days = sorted(int(day) for day in daily.get(treatment, {}))
        missing_models = [
            model_number
            for model_number in land_models.get(treatment, ())
            if not model_humidity_sensors.get(model_number)
        ]
        coverage[treatment] = {
            "sensor_ids": treatment_sensors[treatment],
            "days": len(days),
            "first_day_of_year": days[0] if days else None,
            "last_day_of_year": days[-1] if days else None,
            "unmatched_model_numbers": missing_models,
        }
    return {
        "workbook": str(workbook),
        "workbook_sha256": _file_sha256(workbook),
        "aggregation": "median within sensor/day, then median across sensors/land",
        "valid_range_percent": [0.0, 100.0],
        "rejected_values": rejected_values,
        "comparison_scope": "normalized temporal dynamics only; sensor percent is not converted to root-zone extractable water in mm",
        "coverage": coverage,
        "daily": dict(daily),
    }


def build_2024_field_validation_cache(
    workbook: str | Path,
    weather_csv: str | Path,
    template_json: str | Path,
    destination: str | Path,
    start_day: int = 90,
    end_day: int = 300,
) -> dict:
    """由 SmartFarmDB 原始表构建 2024 年实测外部验证缓存。"""
    from openpyxl import load_workbook

    workbook = Path(workbook)
    weather_csv = Path(weather_csv)
    template_json = Path(template_json)
    with template_json.open("r", encoding="utf-8") as handle:
        template_record = json.load(handle)[0]

    # 论文给的是 4 月 10 日这一日历日期。2024 为闰年，对应 DOY 101。
    paper_sowing_date = date(
        2024,
        PAPER_FIELD_PROTOCOL["sowing_month"],
        PAPER_FIELD_PROTOCOL["sowing_day"],
    )
    paper_planting_day = paper_sowing_date.timetuple().tm_yday
    template_record["Plant_time(How_many_days_of_the_year)"] = paper_planting_day

    book = load_workbook(workbook, read_only=True, data_only=True)
    management: dict[str, list[tuple]] = defaultdict(list)
    management_order: list[str] = []
    for row in book["水肥管理"].iter_rows(min_row=3, values_only=True):
        if not str(row[1] or "").startswith("2024-") or str(row[3]).strip() != "棉花":
            continue
        treatment = str(row[0]).strip()
        if treatment not in management:
            management_order.append(treatment)
        management[treatment].append(row)

    yields: dict[str, list[tuple[str, float]]] = defaultdict(list)
    for row in book["产量信息"].iter_rows(min_row=3, values_only=True):
        if str(row[1]).strip() != "2024" or row[13] in (None, "", "-"):
            continue
        yields[str(row[0]).strip()].append((str(row[2]).strip(), float(row[13])))
    book.close()

    treatments = [name for name in management_order if name in yields]
    if not treatments:
        raise ValueError("未找到同时具有 2024 水肥管理和实测产量的处理")
    for treatment in treatments:
        if len(management[treatment]) != 12:
            raise ValueError(f"{treatment} 的水肥事件不是 12 次")
        if len(yields[treatment]) != 3:
            raise ValueError(f"{treatment} 的测产样方不是 3 个")

    weather_daily, tav, amp = _read_2024_weather(weather_csv)
    weather_static = {
        key: value
        for key, value in template_record["wthear_data"].items()
        if not isinstance(value, dict)
    }
    weather_static["TAV"] = tav
    weather_static["AMP"] = amp
    template_record["wthear_data"] = {**weather_static, **weather_daily}

    static_template = _build_static_template(template_record)
    sample_count = len(treatments) * 3
    days = end_day - start_day + 1
    sequence = np.empty((sample_count, days, len(SEQUENCE_FEATURES)), dtype=np.float32)
    static = np.empty((sample_count, len(static_template)), dtype=np.float32)
    wl_features = np.empty((sample_count, 1997), dtype=np.float32)
    labels = np.empty(sample_count, dtype=np.float32)
    planting_day = np.full(sample_count, paper_planting_day, dtype=np.int16)
    groups = np.full(sample_count, _environment_id(template_record), dtype="U64")
    sample_ids = np.empty(sample_count, dtype="U64")
    treatment_ids = np.empty(sample_count, dtype="U16")

    management_summary = {}
    index = 0
    for treatment in treatments:
        record = dict(template_record)
        events = management[treatment]
        series = {name: {} for name in (
            "irrigattons(mm)", "fertilizers_N(kg/ha)",
            "fertilizers_P(kg/ha)", "fertilizers_K(kg/ha)",
        )}
        for row in events:
            current = date.fromisoformat(str(row[1])[:10])
            code = f"{current.year % 100:02d}{current.timetuple().tm_yday:03d}"
            series["irrigattons(mm)"][code] = _to_float(row[4]) * 1.5
            series["fertilizers_N(kg/ha)"][code] = _to_float(row[5])
            series["fertilizers_P(kg/ha)"][code] = _to_float(row[6])
            series["fertilizers_K(kg/ha)"][code] = _to_float(row[7])
        record.update(series)
        management_summary[treatment] = {
            "events": len(events),
            "irrigation_mm": sum(series["irrigattons(mm)"].values()),
            "N_kg_ha": sum(series["fertilizers_N(kg/ha)"].values()),
            "P2O5_kg_ha": sum(series["fertilizers_P(kg/ha)"].values()),
            "K2O_kg_ha": sum(series["fertilizers_K(kg/ha)"].values()),
        }
        prepared_sequence = _management_sequence(record, start_day, end_day)
        prepared_static = np.asarray([_get_value(record, item) for item in static_template], dtype=np.float32)
        prepared_wl_features = _wl_feature_vector(record)
        for plot, observed_yield in yields[treatment]:
            sequence[index] = prepared_sequence
            static[index] = prepared_static
            wl_features[index] = prepared_wl_features
            labels[index] = observed_yield
            sample_ids[index] = f"2024/{treatment}/{plot}"
            treatment_ids[index] = treatment
            index += 1

    excluded = sorted(set(yields) - set(treatments))
    metadata = {
        "source": "SmartFarmDB 2024 field measurements",
        "workbook": str(workbook),
        "workbook_sha256": _file_sha256(workbook),
        "weather_csv": str(weather_csv),
        "weather_sha256": _file_sha256(weather_csv),
        "template_json": str(template_json),
        "template_sha256": _file_sha256(template_json),
        "records": sample_count,
        "treatments": treatments,
        "excluded_yield_treatments_without_2024_management": excluded,
        "start_day": start_day,
        "end_day": end_day,
        "sequence_features": list(SEQUENCE_FEATURES),
        "static_features": ["/".join(item) for item in static_template],
        "wl_dnn_features": 1997,
        "wl_dnn_leap_year_policy": "drop February 29 to preserve the original 365-day calendar feature layout",
        "management_summary": management_summary,
        "paper_field_protocol": {
            **PAPER_FIELD_PROTOCOL,
            "validation_sowing_date": paper_sowing_date.isoformat(),
            "validation_planting_day_of_year": paper_planting_day,
            "date_transfer_assumption": "apply the paper's April 10 field protocol to the 2024 validation year",
        },
        "unit_conversion": "irrigation 1 m3/mu = 1.5 mm; weather wind km/day / 86.4 = m/s; fertilizer P/K retain the SmartFarmDB P2O5/K2O convention without elemental conversion",
        "fertilizer_semantics": {
            "N": "elemental N",
            "P": "P2O5 workbook convention",
            "K": "K2O workbook convention",
            "conversion_to_elemental": None,
            "evidence": "SmartFarmDB IFManagementInfo and workbook headers explicitly name P2O5/K2O; the training JSON generator copies these columns without conversion",
        },
        "planting_day_source": "Wang et al. (2024), field experiment machine sowing on April 10",
        "planting_day_source_verified": True,
        "planting_day_2024_field_observed": False,
        "planting_day_note": "The paper reports April 10 for the 2023 experiment; the same calendar date is transferred to 2024 by user-specified protocol.",
        "static_input_limit": "initial SH2O/SNH4/SNO3 and DSSAT soil fields inherit the existing 2024data.json template; they are not 2024 treatment-level observations",
        "weather_static": {
            "TAV": tav,
            "AMP": amp,
            "source": "computed from the 366 daily 2024 weather records",
        },
    }
    destination = Path(destination)
    destination.parent.mkdir(parents=True, exist_ok=True)
    np.savez(
        destination,
        sequence=sequence,
        static=static,
        wl_features=wl_features,
        labels=labels,
        planting_day=planting_day,
        groups=groups,
        sample_ids=sample_ids,
        treatment_ids=treatment_ids,
        metadata=np.array(json.dumps(metadata, ensure_ascii=False)),
    )
    return metadata


def build_cache(
    source: str | Path,
    destination: str | Path,
    expected_records: int,
    start_day: int = 90,
    end_day: int = 300,
) -> dict:
    """将原始大 JSON 转成一个可复用的 NumPy 缓存。"""
    iterator = stream_json_array(source)
    first = next(iterator)
    template = _build_static_template(first)
    days = end_day - start_day + 1
    sequence = np.empty((expected_records, days, len(SEQUENCE_FEATURES)), dtype=np.float32)
    static = np.empty((expected_records, len(template)), dtype=np.float32)
    wl_features = np.empty((expected_records, 1997), dtype=np.float32)
    labels = np.empty(expected_records, dtype=np.float32)
    planting_day = np.empty(expected_records, dtype=np.int16)
    groups = np.empty(expected_records, dtype="U64")

    count = 0
    for count, record in enumerate(itertools.chain((first,), iterator), start=1):
        if count > expected_records:
            raise ValueError("实际记录数超过 expected_records")
        index = count - 1
        if tuple(record) != tuple(first):
            raise ValueError(f"第 {count} 条顶层字段顺序与首条不一致")
        sequence[index] = _management_sequence(record, start_day, end_day)
        static[index] = [_get_value(record, item) for item in template]
        wl_features[index] = _wl_feature_vector(record)
        labels[index] = float(record["yield_HWANS(kg/ha)"])
        planting_day[index] = int(record["Plant_time(How_many_days_of_the_year)"])
        groups[index] = _environment_id(record)
    if count != expected_records:
        raise ValueError(f"预期 {expected_records} 条，实际 {count} 条")

    metadata = {
        "source": str(source),
        "records": count,
        "start_day": start_day,
        "end_day": end_day,
        "sequence_features": list(SEQUENCE_FEATURES),
        "static_features": ["/".join(item) for item in template],
        "wl_dnn_features": 1997,
        "wl_dnn_feature_order": "Wang legacy JSON insertion order",
        "unique_groups": int(np.unique(groups).size),
        "zero_yields": int(np.count_nonzero(labels == 0)),
    }
    destination = Path(destination)
    destination.parent.mkdir(parents=True, exist_ok=True)
    np.savez(
        destination,
        sequence=sequence,
        static=static,
        wl_features=wl_features,
        labels=labels,
        planting_day=planting_day,
        groups=groups,
        metadata=np.array(json.dumps(metadata, ensure_ascii=False)),
    )
    return metadata


def group_split(groups: np.ndarray, seed: int = 42) -> dict[str, np.ndarray]:
    unique_groups = np.unique(groups)
    rng = np.random.default_rng(seed)
    rng.shuffle(unique_groups)
    train_end = round(len(unique_groups) * 0.70)
    val_end = train_end + round(len(unique_groups) * 0.15)
    group_sets = {
        "train": unique_groups[:train_end],
        "validation": unique_groups[train_end:val_end],
        "test": unique_groups[val_end:],
    }
    return {
        name: np.flatnonzero(np.isin(groups, selected))
        for name, selected in group_sets.items()
    }


def derive_soil_water_features(
    static: np.ndarray, metadata: dict
) -> tuple[np.ndarray, np.ndarray]:
    """从 DSSAT 土层字段推导有效持水量和初始根区水量，单位均为 mm。"""
    feature_index = {name: index for index, name in enumerate(metadata["static_features"])}
    bottoms = np.asarray([5, 15, 30, 60, 100, 200], dtype=np.float32)
    thickness = np.diff(np.concatenate(([0.0], bottoms)))
    sh2o_depths = np.asarray([5, 10, 20, 30, 50, 70, 100, 130, 160, 190], dtype=np.float32)
    sh2o_paths = [f"record/SH2O/{int(depth)}" for depth in sh2o_depths]
    required = []
    for bottom in bottoms.astype(int):
        required.extend(
            [
                f"soil/SLB_{bottom}/SLLL",
                f"soil/SLB_{bottom}/SDUL",
            ]
        )
    required.extend(sh2o_paths)
    missing = [path for path in required if path not in feature_index]
    if missing:
        raise ValueError(f"静态特征缺少土壤水分推导字段：{missing}")

    lower = np.stack(
        [static[:, feature_index[f"soil/SLB_{int(bottom)}/SLLL"]] for bottom in bottoms],
        axis=1,
    )
    upper = np.stack(
        [static[:, feature_index[f"soil/SLB_{int(bottom)}/SDUL"]] for bottom in bottoms],
        axis=1,
    )
    sh2o = np.stack([static[:, feature_index[path]] for path in sh2o_paths], axis=1)
    sh2o_at_bottom = np.stack(
        [np.interp(bottoms, sh2o_depths, row) for row in sh2o], axis=0
    )
    available = np.maximum(upper - lower, 0.0)
    initial_fraction = np.clip(sh2o_at_bottom, lower, upper) - lower
    capacity = np.sum(available * thickness[None, :] * 10.0, axis=1)
    initial = np.sum(np.maximum(initial_fraction, 0.0) * thickness[None, :] * 10.0, axis=1)
    initial = np.minimum(initial, capacity)
    if not np.isfinite(capacity).all() or not np.isfinite(initial).all():
        raise ValueError("土壤水分推导出现非有限值")
    if np.any(capacity <= 0):
        raise ValueError("土壤有效持水量必须为正")
    return capacity.astype(np.float32), initial.astype(np.float32)
